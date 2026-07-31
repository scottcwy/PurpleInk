from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT = Path("docs/artifacts/宏泰AI智能体-产品需求文档-v1.docx")
FONT = "Microsoft YaHei"
ACCENT = "1F4E79"
MUTED = "666666"
LIGHT = "EAF2F8"


def set_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if italic is not None:
        run.italic = italic


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, fill=None):
    if fill:
        shade(cell, fill)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    set_font(r, 9.5, bold=bold, color="000000")


def set_cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def add_rule(paragraph, color=ACCENT):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "8")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def add_body(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.35
    r = p.add_run(text)
    set_font(r, 10.5)
    return p


def add_step(doc, label, text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.left_indent = Cm(0.74)
    p.paragraph_format.first_line_indent = Cm(-0.74)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.3
    r = p.add_run(label)
    set_font(r, 10.5, bold=True, color=ACCENT)
    r = p.add_run(text)
    set_font(r, 10.5)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    set_font(r, 10.5)
    return p


def add_h1(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_font(r, 16, bold=True, color=ACCENT)
    return p


def add_h2(doc, text):
    p = doc.add_paragraph(style="Heading 2")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_font(r, 12.5, bold=True, color="1F4E79")
    return p


def add_label_paragraph(doc, label, text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.3
    r = p.add_run(label)
    set_font(r, 10.5, bold=True, color=ACCENT)
    r = p.add_run(text)
    set_font(r, 10.5)
    return p


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.35
    for style_name, size, before, after, color in [
        ("Heading 1", 16, 18, 9, ACCENT),
        ("Heading 2", 12.5, 12, 6, "1F4E79"),
    ]:
        style = doc.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def configure_section(section):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.3)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)


def add_header_footer(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("宏泰AI智能体｜产品需求文档（第一版）")
    set_font(r, 8.5, color=MUTED)
    add_rule(p, "B7C9D6")

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("第 ")
    set_font(r, 8.5, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)
    r = p.add_run(" 页")
    set_font(r, 8.5, color=MUTED)


def add_metadata_table(doc):
    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    entries = [
        ("文档名称", "宏泰AI智能体 - 产品需求文档"),
        ("版本", "V1.0（第一版）"),
        ("文档类型", "产品需求文档（PRD）"),
        ("状态", "Demo 阶段需求基线"),
    ]
    for row, (label, value) in zip(table.rows, entries):
        for cell in row.cells:
            set_cell_margins(cell)
        set_cell_text(row.cells[0], label, bold=True, fill=LIGHT)
        set_cell_text(row.cells[1], value)
    return table


def add_contents(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("目录")
    set_font(r, 20, bold=True, color=ACCENT)
    contents = [
        "一、这个软件到底是做什么的",
        "二、账号系统：注册和登录",
        "三、整体页面结构：底部四个按钮",
        "四、“拆解”页面：把别人的爆款视频变成可用的素材",
        "五、“制作”页面：用模板快速做出新视频",
        "六、“素材”页面：所有拆解过的内容和素材都在这里",
        "七、“设置”页面：账号信息和个人资料",
        "八、关于“发布”这件事：全程只用手机，不用电脑",
        "九、这一版先不做的功能",
        "十、技术栈说明",
        "十一、部署与交付说明",
    ]
    for index, item in enumerate(contents, 1):
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.left_indent = Cm(0.3)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.2
        r = p.add_run(f"{index:02d}  {item}")
        set_font(r, 10.5, color="333333")


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0])
    add_header_footer(doc.sections[0])
    doc.core_properties.title = "宏泰AI智能体 - 产品需求文档（第一版）"
    doc.core_properties.subject = "产品需求文档"
    doc.core_properties.author = "PurpleInk"
    doc.core_properties.comments = "已按 PRD 版式完成排版。"

    # Cover
    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("宏泰AI智能体")
    set_font(r, 26, bold=True, color=ACCENT)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(32)
    r = p.add_run("产品需求文档（第一版）")
    set_font(r, 18, color="333333")
    add_metadata_table(doc)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # Contents
    add_contents(doc)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    add_h1(doc, "一、这个软件到底是做什么的")
    add_body(doc, "简单来说，这是一个装在手机里的“内容创作助手”。它主要帮你做一件事：把别人的爆款短视频，变成你自己能用的爆款短视频。")
    add_body(doc, "具体来说，整个使用过程是这样的：你刷到一条很火的短视频，觉得这条视频的套路很好，就把这条视频的链接复制粘贴到我们的软件里。软件会自动帮你把视频里说的话整理成文字，再帮你分析这条视频为什么火（开头怎么抓人、中间怎么讲道理、最后怎么让人心动下单），然后把这个“爆款套路”变成一个可以反复使用的“模板”。")
    add_body(doc, "以后你想拍新视频的时候，不用再从零想创意，只要选一个之前存好的模板，输入你自己的产品或者话题，软件就能帮你把新的文案写出来，再从你自己上传的素材里挑视频片段，自动拼成一条新视频。做好之后，你在手机上点一下“发布”，视频就能直接发到你的短视频账号上，全程不需要用电脑。")
    add_body(doc, "也就是说，这个软件把“看爆款 → 拆解爆款 → 模仿爆款 → 做出新视频 → 发布”这一整套原本需要好几个软件、好几个步骤才能完成的事情，全部整合在一个手机应用里，全程动动手指就能完成。")

    add_h1(doc, "二、账号系统：注册和登录")
    add_body(doc, "用户第一次打开软件，需要先注册一个账号，之后每次打开软件都需要登录。")
    add_body(doc, "为了让用户注册的时候更简单、更快，第一版我们不使用手机验证码这种方式（因为开通短信验证码功能需要额外申请资质、也要花钱，现阶段项目刚起步，暂时不需要）。改成用户直接设置一个账号名和密码就能完成注册，简单快速，不需要等短信、不需要跳转其他软件验证身份。")
    add_body(doc, "后续如果软件用户变多了，觉得有必要增加安全性，我们随时可以在这个基础上加上验证码功能。")

    add_h1(doc, "三、整体页面结构：底部四个按钮")
    add_body(doc, "打开软件之后，屏幕最下面有四个按钮，分别是“拆解”“制作”“素材”“设置”，用户点哪个按钮就切换到对应的页面。这四个页面各自负责一件事，配合起来就是完整的工作流程。")
    for name, desc in [("拆解", "将短视频链接处理为文字稿、爆款结构分析与可复用模板。"), ("制作", "选择模板并输入主题，生成可预览、可发布的新视频。"), ("素材", "管理拆解模板、个人视频和图片素材。"), ("设置", "维护账号信息与生意档案。")]:
        add_label_paragraph(doc, f"{name}：", desc)

    add_h1(doc, "四、“拆解”页面：把别人的爆款视频变成可用的素材")
    add_body(doc, "这是用户打开软件后最常用的页面。页面正中间是一个粘贴框，用户把从短视频软件里复制的链接粘贴进来，点一下“开始拆解”按钮，软件就会按照下面这几步自动处理：")
    add_step(doc, "第一步，取出干净的视频（去水印技术）。", "软件会自动读取这个链接，把视频本身下载下来，并且去掉视频上原本平台加的水印标志，方便用户后续拿这个视频做参考或者素材。")
    add_step(doc, "第二步，把视频里说的话变成文字（语音转文字技术）。", "很多爆款视频的文案不是写在简介里的，而是主播直接讲出来的。软件会自动“听”视频里的声音，把说的话整理成一段完整的文字稿，用户不需要自己一句一句抄。")
    add_step(doc, "第三步，分析这条视频为什么火（智能内容分析技术）。", "软件会把整理出来的文字稿交给背后的智能分析引擎，自动拆解出这条视频的“套路”：开头是怎么一句话抓住人的注意力的、中间是怎么讲道理或者讲痛点的、后面是怎么一步步让人心动的、结尾是怎么引导下单或者关注的。这个拆解结果会用简单易懂的方式展示给用户看，让用户一眼就明白这条视频“火在哪里”。")
    add_step(doc, "第四步，生成一个可以重复使用的模板。", "拆解完之后，软件会把这个“套路”变成一个模板，存起来。以后用户做新视频的时候，可以直接选这个模板，把里面的内容换成自己的产品或者话题，而不用每次都重新想。")
    add_body(doc, "拆解完成后，这条记录会自动保存到“素材”页面里，方便用户以后随时找出来用。")

    add_h1(doc, "五、“制作”页面：用模板快速做出新视频")
    add_body(doc, "这个页面是用户真正开始动手制作自己视频的地方。页面中间有一个输入框，用户在这里输入自己这次想拍的主题（比如“我们家的产品有多安全”），下面可以选择一个之前拆解好、存起来的模板。")
    add_body(doc, "选好主题和模板之后，用户点一下“一键制作”，软件会自动做下面几件事：")
    add_step(doc, "第一步，生成属于用户自己的新文案（智能文案生成技术）。", "软件会按照选好的模板结构，结合用户输入的主题，自动写出一段全新的文案，风格套路和原来的爆款视频类似，但内容完全是围绕用户自己的产品或者话题来写的。")
    add_step(doc, "第二步，自动挑选合适的视频片段（智能素材匹配技术）。", "软件会从用户自己上传过的视频素材库里，按照新文案的节奏，自动挑选合适的片段。")
    add_step(doc, "第三步，自动拼接成一条完整的视频（智能视频合成技术）。", "把挑好的素材片段和文案自动拼接、配上字幕，做成一条完整的、可以直接发布的短视频。")
    add_body(doc, "做好之后，用户可以在手机上直接预览这条新视频，觉得满意的话，可以直接点“发布”，把视频发到自己的短视频账号上（这个发布过程完全在手机上完成，不需要打开电脑，具体见下面第八部分）。")

    add_h1(doc, "六、“素材”页面：所有拆解过的内容和素材都在这里")
    add_body(doc, "这个页面就是用户的个人素材仓库，主要分成两大块：")
    add_h2(doc, "1. 拆解模板库")
    add_body(doc, "所有用户之前拆解过的爆款视频，都会按照拆解时间自动整理在这里，用户可以随时点开看当时的文字稿和拆解结果，也可以在这里直接选择某个模板去“制作”页面使用。")
    add_h2(doc, "2. 我的素材")
    add_body(doc, "用户自己拍的视频、图片素材，可以上传到这里，并且可以自己建文件夹分类整理（比如按门店、按产品、按场景分类），这样在“制作”页面自动挑素材的时候，才能更准确地找到合适的片段。")

    add_h1(doc, "七、“设置”页面：账号信息和个人资料")
    add_body(doc, "这个页面主要放两类内容：")
    add_h2(doc, "1. 账号信息")
    add_body(doc, "用户可以在这里查看和修改自己的账号资料、修改密码、退出登录。")
    add_h2(doc, "2. 我的生意档案（用于让软件更懂你）")
    add_body(doc, "用户可以在这里填写一些关于自己生意的信息，比如做什么行业、卖什么产品、门店在哪个城市、自己的人设是什么样的。这些信息填得越详细，软件在“制作”页面帮用户写文案的时候，就能写得越贴合用户自己的实际情况，而不是千篇一律的通用内容。")

    add_h1(doc, "八、关于“发布”这件事：全程只用手机，不用电脑")
    add_body(doc, "这是这个软件的一个重要卖点：用户做完视频之后，不需要打开电脑、不需要登录任何网页后台，直接在手机软件里点“发布”，视频就能自动发到自己的短视频账号上。")
    add_body(doc, "具体是怎么做到的：第一次使用发布功能的时候，软件会在手机屏幕上弹出一个二维码，用户拿自己的短视频账号扫一下这个码，完成一次授权（就像平时用手机扫码登录网页版一样），后续每次点“发布”，软件后台会自动帮用户把视频发布上去，用户全程只需要在手机上操作，不涉及电脑（自动发布技术）。")

    add_h1(doc, "九、这一版先不做的功能")
    add_body(doc, "为了让第一版尽快做出来、尽快能用，下面这些功能这一版先不做，留到以后有需要再加：")
    add_bullet(doc, "虚拟数字人出镜、克隆自己的声音说话，这一版不做。")
    add_bullet(doc, "一键复制别人视频的画面风格来拍同款，这一版不做。")
    add_bullet(doc, "用碰一碰手机的方式让别人帮忙转发视频，这一版不做。")
    add_bullet(doc, "持续 monitoring 某个同行账号、自动追踪他每次发新视频，这一版先不做，用户可以自己手动把想参考的视频链接粘贴进来拆解。")

    add_h1(doc, "十、技术栈说明")
    add_body(doc, "本项目采用现代化全栈架构进行开发，前端以 Next.js 为核心框架，结合 TypeScript 实现类型安全与工程化开发，页面交互基于 React 生态构建。样式层可采用 Tailwind CSS 或同类组件化方案，以提升开发效率与界面一致性。")
    add_body(doc, "后端逻辑基于 Node.js 运行环境实现，结合服务端 API 接口承载用户认证、内容处理、任务调度、模板管理与结果回传等核心能力。数据层在 Demo 阶段优先采用轻量化存储方案，便于快速迭代与功能验证；后续正式版本可平滑迁移至云数据库或自建数据库服务。")
    add_body(doc, "在 AI 能力层面，系统将通过统一的模型接入层调用多个大模型服务，支持文本生成、内容拆解、结构分析、模板改写、任务编排等功能。所有模型调用均通过环境变量或受控配置方式管理，确保项目具备较好的可维护性与扩展性。")

    add_h1(doc, "十一、部署与交付说明")
    add_body(doc, "当前项目 Demo 阶段将部署在腾讯云 EdgeOne Pages 平台，利用其前后端一体化部署能力完成快速验证。Demo 的部署有效期为三个月，主要用于功能演示、流程验证与客户确认。")
    add_body(doc, "在 Demo 交付完成后，项目将进入正式部署与客户接管阶段。届时可根据客户实际使用情况，逐步引导客户完成以下事项：")
    for item in [
        "开通并配置客户自己的 AI 模型服务，并获取对应的 API Key。",
        "将当前项目迁移部署到客户自己的腾讯云 EdgeOne Pages 环境，或迁移至其自有云服务器环境。",
        "根据客户实际业务规模，选择适合的数据库、存储与任务执行方案。",
        "完成域名、环境变量、接口权限与上线配置的正式交接。",
    ]:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        r = p.add_run(item)
        set_font(r, 10.5)
    add_body(doc, "当前阶段不对正式交付的部署方式做强制锁定，优先目标是先完成 Demo 验证、核心流程跑通以及客户需求对齐。正式交付方案可在 Demo 验收后再进一步细化。")

    doc.save(OUT)
    print(OUT.resolve())


if __name__ == "__main__":
    main()
